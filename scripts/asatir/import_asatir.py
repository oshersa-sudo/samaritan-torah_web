# -*- coding: utf-8 -*-
"""Import ספר האסאטיר (the Samaritan "Asatir" / Book of the Secrets of Moses) as
a Samaritan Library unit AND as a new "ממקור שומרון" source.

The book is a running para-biblical chronicle from Adam to the end of days, so
unlike the other sources it is not a commentary keyed to verses — it RETELLS the
Torah's episodes. Two signals therefore tie a passage to the Torah:

  1. NARRATIVE CORRESPONDENCE (method 'ref') — the hand-authored LINKS map below.
     Each entry is "chapter.paragraph" -> the Torah verses that paragraph retells,
     assigned by reading the passage against the Torah text. This is the primary
     signal: the Asatir paraphrases rather than cites, so most of its connection
     to a verse leaves no quotable trace.
  2. QUOTED SCRIPTURE (method 'quote') — the book does quote the Torah outright
     inside quotation marks ("זה ינחמנו ממעשינו ומעצבון ידינו"), and those are
     matched on the consonantal skeleton by scripts/link_source_to_verses.py, the
     same matcher the אם בחקותי commentary uses.

Tables are additive only: asatir_sections / asatir_verse_links. The Torah text
itself is never touched.

Usage:  py -3 scripts/asatir/import_asatir.py            # dry run: parse+validate
        py -3 scripts/asatir/import_asatir.py --apply
"""
import os
import re
import io
import sys
import shutil
import sqlite3
import datetime

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
_HERE = os.path.dirname(os.path.abspath(__file__))
_ROOT = os.path.dirname(os.path.dirname(_HERE))
sys.path.insert(0, os.path.join(_ROOT, 'scripts'))

from link_source_to_verses import VerseIndex, QUOTE_RE   # noqa: E402

APPLY = '--apply' in sys.argv
DOCX = os.path.join(_ROOT, 'אסאטיר.docx')
DB = os.path.join(_ROOT, 'data', 'torah.db')

G, E, L, N, D = 'בראשית', 'שמות', 'ויקרא', 'במדבר', 'דברים'
HEB_NUM = ['', 'א', 'ב', 'ג', 'ד', 'ה', 'ו', 'ז', 'ח', 'ט', 'י', 'יא', 'יב',
           'יג', 'יד', 'טו', 'טז', 'יז', 'יח', 'יט', 'כ']


def gv(book, *cvs):
    return [(book, c, v) for c, v in cvs]


# ── section ("chapter.paragraph") -> the Torah verses that paragraph retells ──
# Paragraph numbering is the docx's own paragraph order inside each H2 chapter;
# scripts/asatir/import_asatir.py --dry prints the same numbers, so the map can
# be re-checked passage by passage. A paragraph with no Torah counterpart (the
# apocalyptic "יקום נשיא" oracles, the closing doxology) is deliberately absent.
LINKS = {
    # ── פרק א · אדם קין הבל ──
    '1.1':  gv(G, (4, 1), (4, 2), (4, 17)),
    '1.2':  gv(G, (4, 3), (4, 4)),
    '1.3':  gv(G, (4, 3)),
    '1.4':  gv(G, (4, 5)),
    '1.5':  gv(G, (4, 5)),
    '1.6':  gv(G, (4, 4), (4, 5)),
    '1.7':  gv(G, (4, 8)),
    '1.8':  gv(G, (4, 8)),
    '1.9':  gv(G, (3, 6), (4, 10), (4, 11)),
    '1.10': gv(G, (1, 27), (2, 15), (3, 23), (3, 24)),
    '1.11': gv(G, (4, 25), (5, 3)),
    # ── פרק ב · אלה תולדות אדם ──
    '2.1':  gv(G, (4, 16), (4, 17)),
    '2.2':  gv(G, (4, 26), (5, 6), (5, 9), (5, 12), (5, 15)),
    '2.3':  gv(G, (5, 18)),
    '2.4':  gv(G, (5, 21), (5, 22)),
    '2.5':  gv(G, (4, 18), (4, 23), (4, 24), (6, 11)),
    '2.6':  gv(G, (5, 22), (5, 25), (5, 28)),
    '2.7':  gv(G, (5, 29)),
    '2.8':  gv(G, (4, 20), (4, 21), (4, 22)),
    # ── פרק ג · חנוך, נח, ואחידן בן ברד ──
    '3.1':  gv(G, (5, 5), (5, 8), (5, 11), (5, 14), (5, 17), (5, 20), (5, 23)),
    '3.2':  gv(G, (5, 24)),
    '3.3':  gv(G, (5, 24)),
    '3.4':  gv(G, (5, 24)),
    '3.5':  gv(G, (5, 27)),
    '3.6':  gv(G, (5, 31)),
    '3.7':  gv(G, (7, 6)),
    '3.8':  gv(G, (2, 17), (5, 5)),
    '3.9':  gv(G, (5, 1), (23, 9)),
    '3.10': gv(G, (6, 1)),
    '3.11': gv(G, (6, 5), (6, 11)),
    # ── פרק ד · אחידן וגפנה, נח איש צדיק ──
    '4.1':  gv(G, (4, 22)),
    '4.2':  gv(G, (6, 11), (6, 12)),
    '4.3':  gv(E, (4, 20)),
    '4.4':  gv(G, (6, 12)),
    '4.5':  gv(G, (6, 12)),
    '4.6':  gv(G, (5, 32), (6, 13)),
    '4.7':  gv(G, (6, 18), (7, 7)),
    '4.8':  gv(G, (6, 14), (7, 11)),
    '4.9':  gv(G, (8, 3), (8, 13)),
    '4.10': gv(G, (8, 18), (8, 20), (8, 21), (9, 13)),
    '4.11': gv(G, (9, 9)) + gv(D, (6, 4)),
    # ── פרק ה · אלה תולדות נח ──
    '5.1':  gv(G, (10, 1), (10, 32)),
    '5.2':  gv(G, (10, 22)),
    '5.3':  gv(G, (10, 2)),
    '5.4':  gv(G, (10, 6)),
    '5.5':  gv(G, (9, 28)),
    '5.6':  gv(G, (10, 32)),
    '5.7':  gv(G, (9, 29)),
    '5.8':  gv(G, (10, 2), (10, 22)),
    '5.9':  gv(G, (10, 8), (10, 10), (11, 2)),
    '5.10': gv(G, (9, 29), (10, 21)),
    '5.11': gv(G, (9, 29), (10, 11), (10, 12)),
    # ── פרק ו · מעשה בבל וממלכת נמרוד ──
    '6.1':  gv(G, (11, 2), (11, 3), (11, 4)),
    '6.2':  gv(G, (11, 7), (11, 8), (11, 9)),
    '6.3':  gv(G, (11, 9)),
    '6.5':  gv(G, (15, 7), (26, 1), (26, 5)),
    '6.6':  gv(G, (10, 13), (10, 14), (10, 30)),
    '6.7':  gv(G, (11, 10)) + gv(E, (1, 16)),
    '6.8':  gv(G, (11, 26)),
    # ── פרק ז · לך לך מארצך וממולדך ──
    '7.1':  gv(G, (11, 26)),
    '7.2':  gv(G, (11, 28)),
    '7.3':  gv(G, (11, 31), (12, 1), (12, 5), (12, 6), (14, 1)),
    '7.4':  gv(G, (12, 7), (12, 8), (12, 10)),
    '7.5':  gv(G, (12, 10), (12, 11)),
    '7.6':  gv(G, (12, 14), (12, 15), (12, 16), (12, 17)),
    '7.7':  gv(G, (12, 17), (12, 18), (12, 19)),
    # ── פרק ח · מלחמת המלכים ──
    '8.1':  gv(G, (12, 17), (12, 20)),
    '8.2':  gv(G, (12, 19), (12, 20)),
    '8.3':  gv(G, (13, 1), (13, 3), (13, 4)),
    '8.4':  gv(G, (13, 11), (14, 1), (14, 4)),
    '8.5':  gv(G, (14, 4), (14, 5), (14, 12), (14, 13)),
    '8.6':  gv(G, (14, 14), (14, 15), (14, 16)),
    '8.7':  gv(G, (14, 18), (14, 19), (14, 20)),
    '8.8':  gv(G, (14, 21), (14, 22), (14, 23)),
    # ── פרק ט · אלה תולדת אברהם ──
    '9.1':  gv(G, (15, 1), (15, 5), (15, 6)),
    '9.2':  gv(G, (15, 18)),
    '9.3':  gv(G, (17, 1), (17, 24), (19, 24), (21, 2)),
    '9.4':  gv(G, (15, 18), (25, 8), (25, 18)),
    '9.5':  gv(G, (16, 12), (28, 9), (36, 8)),
    '9.6':  gv(G, (25, 1), (36, 2), (36, 32), (36, 33)),
    '9.7':  gv(G, (10, 25), (10, 26)),
    '9.8':  gv(G, (36, 34), (36, 35), (36, 36), (36, 37), (36, 38), (36, 39)),
    '9.9':  gv(G, (36, 31)),
    '9.10': gv(G, (28, 10), (37, 2)),
    '9.11': gv(E, (1, 8)),
    # ── פרק י · וילך איש מבית לוי ──
    '10.1': gv(E, (1, 6), (1, 8), (1, 9), (1, 10)),
    '10.2': gv(E, (2, 1), (2, 2)),
    '10.3': gv(E, (1, 7), (1, 15), (1, 16), (1, 17), (1, 22)),
    '10.4': gv(E, (1, 22)),
    '10.5': gv(E, (2, 2)),
    '10.6': gv(E, (2, 3), (2, 4), (2, 5)),
    # ── פרק יא · בימים הרבים ההם ──
    '11.1': gv(E, (2, 5), (2, 6)),
    '11.2': gv(E, (2, 6)),
    '11.3': gv(E, (2, 4)),
    '11.4': gv(E, (2, 7), (2, 8), (2, 9), (2, 10)),
    '11.5': gv(E, (2, 11), (2, 12)),
    '11.6': gv(E, (2, 13), (2, 14), (2, 15), (2, 23)),
    '11.7': gv(E, (2, 24), (3, 1), (3, 6)),
    '11.8': gv(E, (3, 12), (4, 14), (4, 27), (4, 30), (12, 41), (14, 21),
               (15, 25), (17, 13)),
    # ── פרק יב · מעמד הר סיני ──
    '12.1':  gv(G, (1, 5)) + gv(E, (19, 16)),
    '12.2':  gv(G, (1, 1), (1, 3)),
    '12.3':  gv(E, (19, 16), (20, 18)),
    '12.4':  gv(D, (32, 35)),
    '12.5':  gv(E, (40, 17)) + gv(N, (10, 12), (20, 1), (20, 28)),
    '12.6':  gv(N, (21, 1), (21, 2), (21, 3)),
    '12.7':  gv(N, (22, 4), (22, 5), (25, 3)),
    '12.8':  gv(N, (22, 5), (22, 6), (22, 7)),
    '12.9':  gv(N, (23, 4)),
    '12.10': gv(G, (24, 10)) + gv(N, (22, 22), (23, 7)),
    # ── פרק יג · מקללה לברכה ──
    '13.1':  gv(N, (22, 12), (23, 10)),
    '13.2':  gv(N, (23, 5)),
    '13.3':  gv(N, (23, 21), (24, 17)),
    '13.4':  gv(N, (23, 23)),
    '13.5':  gv(N, (24, 18), (24, 19)),
    '13.6':  gv(N, (23, 21), (24, 17)) + gv(D, (18, 18)),
    '13.7':  gv(N, (23, 21)),
    '13.8':  gv(N, (24, 14), (31, 16)),
    '13.9':  gv(N, (25, 1), (25, 2)),
    '13.10': gv(N, (2, 10), (2, 12), (2, 14), (25, 15)),
    '13.11': gv(N, (25, 3), (25, 4), (25, 5)),
    '13.12': gv(N, (25, 6), (25, 8), (25, 9)),
    '13.13': gv(N, (25, 7), (25, 8), (25, 11), (25, 13)),
    # ── פרק יד · מלחמות ישראל במדין ──
    '14.1':  gv(N, (25, 8)),
    '14.2':  gv(N, (25, 12), (25, 13)),
    '14.3':  gv(G, (34, 2)) + gv(N, (25, 1)),
    '14.4':  gv(N, (31, 2), (31, 5), (31, 6)),
    '14.5':  gv(G, (49, 17), (49, 19)) + gv(N, (31, 7)),
    '14.6':  gv(N, (31, 8)),
    '14.7':  gv(N, (31, 8), (35, 19)),
    '14.8':  gv(N, (31, 17), (31, 49), (31, 50)) + gv(D, (23, 4)),
    '14.9':  gv(N, (27, 18), (27, 19), (27, 22)),
    '14.10': gv(N, (27, 18)),
    # ── פרק טו · ויכתב משה את התורה ──
    '15.1':  gv(E, (18, 3), (18, 4)),
    '15.2':  gv(N, (27, 12), (34, 3), (34, 4), (34, 5)) + gv(D, (1, 5), (32, 49)),
    '15.3':  gv(D, (3, 25), (3, 27), (34, 1), (34, 4)),
    '15.4':  gv(D, (4, 25), (31, 20), (31, 29)),
    '15.5':  gv(D, (33, 12)),
    '15.6':  gv(D, (33, 7)),
    # ── פרק טז · אחרית הימים ──
    '16.1':  gv(G, (38, 24)) + gv(N, (24, 14)) + gv(D, (21, 18)),
    '16.2':  gv(D, (4, 27), (28, 64)),
    '16.3':  gv(D, (4, 30)),
    '16.4':  gv(D, (30, 5)),
    '16.5':  gv(L, (25, 10), (25, 11)) + gv(D, (11, 29)),
    '16.6':  gv(D, (18, 18)),
    '16.28': gv(D, (33, 21)),
    '16.29': gv(N, (24, 19)),
    '16.30': gv(E, (4, 17)) + gv(D, (18, 18)),
}


def parse(path):
    """The docx: H1 = book/part title, H2 = chapter, each body paragraph = one unit."""
    import docx
    doc = docx.Document(path)
    chaps = []
    cur = None
    for p in doc.paragraphs:
        sty = (p.style.name if p.style else '') or ''
        txt = re.sub(r'\s+', ' ', p.text).strip()
        if sty == 'Heading 1':
            continue                      # "ספר האסאטיר" / "תעתיק הספר"
        if sty == 'Heading 2':
            if not txt:                   # the file has a few empty H2 placeholders
                continue
            cur = {'title': txt, 'paras': []}
            chaps.append(cur)
        elif cur is not None and txt:
            cur['paras'].append(txt)
    return chaps


def main():
    conn = sqlite3.connect(DB, timeout=60)
    conn.row_factory = sqlite3.Row
    vidx = {}
    for r in conn.execute("""SELECT v.id vid, b.name bk, c.number ch, v.number vn
                             FROM verses v JOIN chapters c ON c.id=v.chapter_id
                             JOIN books b ON b.id=c.book_id"""):
        if str(r['vn']).isdigit():
            vidx.setdefault((r['bk'], r['ch'], int(r['vn'])), r['vid'])

    chaps = parse(DOCX)
    idx = VerseIndex(DB)

    built = []
    problems = []
    mapped_keys = set()
    for ci, c in enumerate(chaps, 1):
        for pi, text in enumerate(c['paras'], 1):
            key = '%d.%d' % (ci, pi)
            links = {}                     # verse_id -> (method, shown)
            for (bk, ch, vn) in LINKS.get(key, []):
                mapped_keys.add(key)
                vid = vidx.get((bk, ch, vn))
                if vid:
                    links[vid] = ('ref', '%s %d:%d' % (bk, ch, vn))
                else:
                    problems.append((key, 'verse not in DB: %s %d:%d' % (bk, ch, vn)))
            # quoted scripture — additive; never overwrites a narrative link
            for m in QUOTE_RE.finditer(text):
                q = m.group(1)
                hits = idx.by_quote(q)
                if len(hits) > 6:          # too generic a phrase to pin anywhere
                    continue
                for vid in hits:
                    links.setdefault(vid, ('quote', q[:60]))
            built.append({'chap': ci, 'chap_title': c['title'], 'para': pi,
                          'ref': '%s,%s' % (HEB_NUM[ci], pi),
                          'text': text, 'links': links})

    unknown = sorted(set(LINKS) - mapped_keys, key=lambda k: [int(x) for x in k.split('.')])
    for k in unknown:
        problems.append((k, 'LINKS key has no matching paragraph in the docx'))

    nref = sum(sum(1 for m, _ in b['links'].values() if m == 'ref') for b in built)
    nq = sum(sum(1 for m, _ in b['links'].values() if m == 'quote') for b in built)
    linked = sum(1 for b in built if b['links'])
    print('chapters: %d   sections: %d   sections with links: %d' % (len(chaps), len(built), linked))
    print('verse links: %d narrative (ref) + %d quoted = %d' % (nref, nq, nref + nq))
    print('distinct Torah verses reached: %d'
          % len({v for b in built for v in b['links']}))
    print('validation issues (%d):' % len(problems))
    for k, m in problems:
        print('  [%s] %s' % (k, m))
    print('\nquoted-scripture matches:')
    for b in built:
        for vid, (meth, shown) in b['links'].items():
            if meth == 'quote':
                print('  [%s] %s' % (b['ref'], shown))

    if not APPLY:
        print('\n[dry-run] re-run with --apply to write')
        conn.close()
        return

    bak = DB + '.bak_asatir_' + datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
    shutil.copy2(DB, bak)
    print('backed up ->', os.path.basename(bak))
    cu = conn.cursor()
    cu.execute('DROP TABLE IF EXISTS asatir_verse_links')
    cu.execute('DROP TABLE IF EXISTS asatir_sections')
    cu.execute('''CREATE TABLE asatir_sections (
        id INTEGER PRIMARY KEY, chap INTEGER, chap_title TEXT,
        para INTEGER, ref TEXT, ord INTEGER, text TEXT)''')
    cu.execute('''CREATE TABLE asatir_verse_links (
        id INTEGER PRIMARY KEY, verse_id INTEGER, section_id INTEGER,
        method TEXT, shown TEXT)''')
    for i, b in enumerate(built):
        cu.execute('INSERT INTO asatir_sections (chap, chap_title, para, ref, ord, text) '
                   'VALUES (?,?,?,?,?,?)',
                   (b['chap'], b['chap_title'], b['para'], b['ref'], i, b['text']))
        sid = cu.lastrowid
        for vid, (meth, shown) in b['links'].items():
            cu.execute('INSERT INTO asatir_verse_links (verse_id, section_id, method, shown) '
                       'VALUES (?,?,?,?)', (vid, sid, meth, shown))
    cu.execute('CREATE INDEX ix_asatir_verse ON asatir_verse_links (verse_id)')
    cu.execute('CREATE INDEX ix_asatir_sect ON asatir_verse_links (section_id)')
    conn.commit()
    ns = conn.execute('SELECT COUNT(*) FROM asatir_sections').fetchone()[0]
    nl = conn.execute('SELECT COUNT(*) FROM asatir_verse_links').fetchone()[0]
    nv = conn.execute('SELECT COUNT(DISTINCT verse_id) FROM asatir_verse_links').fetchone()[0]
    print('APPLIED: %d sections, %d links across %d verses' % (ns, nl, nv))
    conn.close()


if __name__ == '__main__':
    main()
